#!/usr/bin/env python3
"""Validate the BYD workbook and build the browser data file."""

from __future__ import annotations

import argparse
import json
import math
import statistics
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPECTED_HEADERS = [
    "日期",
    "开盘",
    "收盘",
    "最高",
    "最低",
    "成交量",
    "成交额",
    "振幅",
    "涨跌幅",
    "涨跌额",
    "换手率",
]

EXPECTED_ROWS = 242


def parse_args() -> argparse.Namespace:
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--input",
        type=Path,
        default=root / "data" / "source" / "比亚迪后复权历史行情数据.xlsx",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=root / "data" / "byd_2023.js",
    )
    return parser.parse_args()


def require_number(value: Any, row_number: int, column: str) -> float:
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        raise ValueError(f"row {row_number}: {column} is not numeric: {value!r}")
    return float(value)


def require_date(value: Any, row_number: int) -> datetime:
    if not isinstance(value, datetime):
        raise ValueError(f"row {row_number}: date is not a datetime: {value!r}")
    return value


def build_payload(input_path: Path) -> dict[str, Any]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise ValueError(f"expected one worksheet, found {workbook.sheetnames!r}")

    worksheet = workbook[workbook.sheetnames[0]]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("workbook is empty")

    headers = list(rows[0])
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"unexpected headers: {headers!r}")
    if len(rows) - 1 != EXPECTED_ROWS:
        raise ValueError(
            f"expected {EXPECTED_ROWS} data rows, found {len(rows) - 1}"
        )

    records: list[dict[str, Any]] = []
    seen_dates: set[str] = set()

    for row_number, row in enumerate(rows[1:], start=2):
        if len(row) != len(EXPECTED_HEADERS):
            raise ValueError(
                f"row {row_number}: expected {len(EXPECTED_HEADERS)} columns, "
                f"found {len(row)}"
            )

        date_value = require_date(row[0], row_number).date()
        date_iso = date_value.isoformat()
        if date_iso in seen_dates:
            raise ValueError(f"duplicate trading date: {date_iso}")
        seen_dates.add(date_iso)

        record = {
            "date": date_iso,
            "open": require_number(row[1], row_number, "开盘"),
            "close": require_number(row[2], row_number, "收盘"),
            "high": require_number(row[3], row_number, "最高"),
            "low": require_number(row[4], row_number, "最低"),
            "volume": require_number(row[5], row_number, "成交量"),
            "amount": require_number(row[6], row_number, "成交额"),
            "amplitudePct": require_number(row[7], row_number, "振幅"),
            "changePct": require_number(row[8], row_number, "涨跌幅"),
            "change": require_number(row[9], row_number, "涨跌额"),
            "turnoverPct": require_number(row[10], row_number, "换手率"),
        }
        if record["high"] < record["low"]:
            raise ValueError(f"row {row_number}: high is below low")
        records.append(record)

    closes = [record["close"] for record in records]
    highs = [record["high"] for record in records]
    lows = [record["low"] for record in records]
    volumes = [record["volume"] for record in records]
    amounts = [record["amount"] for record in records]
    changes = [record["changePct"] for record in records]
    amplitudes = [record["amplitudePct"] for record in records]

    max_high_index = max(range(len(records)), key=lambda index: highs[index])
    min_low_index = min(range(len(records)), key=lambda index: lows[index])
    period_change = closes[-1] - closes[0]
    period_return_pct = period_change / closes[0] * 100

    stats = {
        "recordCount": len(records),
        "startDate": records[0]["date"],
        "endDate": records[-1]["date"],
        "firstClose": closes[0],
        "lastClose": closes[-1],
        "periodChange": period_change,
        "periodReturnPct": period_return_pct,
        "closeMean": statistics.mean(closes),
        "closeMedian": statistics.median(closes),
        "closeStd": statistics.stdev(closes),
        "highMax": max(highs),
        "highDate": records[max_high_index]["date"],
        "lowMin": min(lows),
        "lowDate": records[min_low_index]["date"],
        "amplitudeMeanPct": statistics.mean(amplitudes),
        "dailyChangeStdPct": statistics.stdev(changes),
        "annualizedVolatilityPct": statistics.stdev(changes)
        * math.sqrt(EXPECTED_ROWS),
        "volumeMean": statistics.mean(volumes),
        "amountMean": statistics.mean(amounts),
        "amountTotal": sum(amounts),
        "upDays": sum(value > 0 for value in changes),
        "downDays": sum(value < 0 for value in changes),
        "flatDays": sum(value == 0 for value in changes),
    }

    stats = {
        key: round(value, 4) if isinstance(value, float) else value
        for key, value in stats.items()
    }

    return {
        "meta": {
            "stockName": "比亚迪",
            "stockCode": "002594",
            "market": "深交所",
            "adjustment": "后复权",
            "sourceFile": input_path.name,
            "sourceSheet": worksheet.title,
            "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
            "units": {
                "price": "元",
                "volume": "手",
                "amount": "元",
                "percentage": "%",
            },
        },
        "stats": stats,
        "rows": records,
    }


def main() -> None:
    args = parse_args()
    payload = build_payload(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    serialized = json.dumps(payload, ensure_ascii=False, indent=2)
    args.output.write_text(
        "/* Generated by scripts/build_data.py. Do not edit by hand. */\n"
        f"window.BYD_DATA = {serialized};\n",
        encoding="utf-8",
    )
    stats = payload["stats"]
    print(
        "Generated "
        f"{args.output} with {stats['recordCount']} rows, "
        f"{stats['startDate']} to {stats['endDate']}."
    )


if __name__ == "__main__":
    main()
